"""Structural checks on the generated workbook. Exits non-zero on failure.

Run:  python3 build_workbook.py && python3 verify_workbook.py

This reads the .xlsx that was actually written rather than the builder's
intentions, which is the whole point: the previous workbook passed its own
checks while shipping 13 untitled charts, every chart anchored over its own
caption, and 201 formula cells with no cached value. Each of those is now a
check below, named for the defect it exists to catch.
"""

import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import build_workbook as B
from build_workbook import CONTENTS, INDEX, OUT
from dataset import OBS
from sources import REFERENCE_ONLY, SOURCES

FAILURES = []
CHECKS = [0]


def check(condition, message):
    CHECKS[0] += 1
    if not condition:
        FAILURES.append(message)


# The sheet names report/render.py and report/draft/appendix/A_regression.md
# cite. The rebuild kept the report untouched, so losing one of these would
# leave a caption pointing at a tab that does not exist.
CITED_BY_REPORT = [
    "02_MASTER", "09_POLICY", "F1_BRANCHES", "F2_PAYMENTS",
    "F6_ADOPT_BENEFIT", "F11_EWASTE", "F12_REACH",
]


def cell_extent(anchor):
    """The (row0, col0, row1, col1) block a drawing covers, 0-indexed."""
    frm = anchor._from
    to = getattr(anchor, "to", None)
    if to is not None:
        return frm.row, frm.col, to.row, to.col
    # oneCellAnchor: fall back to the nominal chart box, 760x460 px, against
    # Excel's default 64px column and 20px row.
    return frm.row, frm.col, frm.row + 23, frm.col + 12


def main():
    if not os.path.exists(OUT):
        print("workbook not built; run build_workbook.py first")
        return 1

    wb = openpyxl.load_workbook(OUT)
    wbv = openpyxl.load_workbook(OUT, data_only=True)
    names = [ws.title for ws in wb.worksheets]

    # -- 1. contents list matches the real tabs, in order --------------------
    listed = [n for n, _ in CONTENTS]
    check(names == listed,
          f"00_COVER contents do not match the tabs.\n  tabs:   {names}\n"
          f"  listed: {listed}")

    cover = wb["00_COVER"]
    on_cover = [cover.cell(row=r, column=1).value
                for r in range(1, cover.max_row + 1)]
    for name in names:
        if name == "00_COVER":
            continue
        check(name in on_cover, f"{name} is a tab but is not listed on 00_COVER")

    # -- 2. every cited sheet still exists -----------------------------------
    for name in CITED_BY_REPORT:
        check(name in names,
              f"{name} is cited by the report but is not a sheet in the workbook")

    # -- 3. master integrity --------------------------------------------------
    master = wb["02_MASTER"]
    headers = [c.value for c in master[1]]
    check(headers == B.MASTER_HEADERS, f"02_MASTER headers changed: {headers}")

    src_sheet = wb["03_SOURCES"]
    src_ids = {src_sheet.cell(row=r, column=1).value
               for r in range(5, src_sheet.max_row + 1)}
    src_ids.discard(None)

    rows = 0
    for r in range(2, master.max_row + 1):
        code = master.cell(row=r, column=B.M_CODE + 1).value
        if code is None:
            continue
        rows += 1
        unit = master.cell(row=r, column=B.M_UNIT + 1).value
        denom = master.cell(row=r, column=B.M_DENOM + 1).value
        sid = master.cell(row=r, column=B.M_SRC + 1).value
        check(bool(unit), f"02_MASTER row {r} ({code}) has no unit")
        check(bool(denom), f"02_MASTER row {r} ({code}) has no denominator")
        check(sid in src_ids,
              f"02_MASTER row {r} ({code}) cites source {sid!r}, "
              f"which is not in 03_SOURCES")
    check(rows == len(OBS),
          f"02_MASTER holds {rows} observations, dataset has {len(OBS)}")

    # -- 4. every source is usable -------------------------------------------
    for r in range(5, src_sheet.max_row + 1):
        sid = src_sheet.cell(row=r, column=1).value
        if sid is None:
            continue
        # A guest lecture has no URL and no retrieval date. Those sources are
        # listed in REFERENCE_ONLY and are held to the reference alone.
        if sid not in REFERENCE_ONLY:
            check(bool(src_sheet.cell(row=r, column=5).value),
                  f"03_SOURCES {sid} has no access date")
            check(bool(src_sheet.cell(row=r, column=6).value),
                  f"03_SOURCES {sid} has no URL")
        check(bool(src_sheet.cell(row=r, column=7).value),
              f"03_SOURCES {sid} has no Harvard reference")

    # -- 5. every formula carries a cached value -----------------------------
    #    The defect that made the previous workbook show blanks and empty
    #    charts in any viewer without a calculation engine.
    formulas = 0
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    cached = wsv[c.coordinate].value
                    check(cached is not None,
                          f"{ws.title}!{c.coordinate} is a formula with no "
                          f"cached value")
    check(formulas > 0, "no formulas found; the workbook should be formula-driven")

    # -- 6. cached values agree with the dataset ------------------------------
    #    A SUMIFS pointing at the wrong observation would still produce a
    #    number, so check the number rather than the formula's shape.
    checked_values = 0
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                f = c.value
                if not f.startswith("=SUMIFS(") or f.count("SUMIFS") != 1:
                    continue      # a composite expression; covered by check 5
                try:
                    code = f.split('"')[1]
                    year = int(f.rsplit(",", 1)[1].rstrip(")"))
                except (IndexError, ValueError):
                    continue
                expected = INDEX.get((code, year), (None,) * 5)[4]
                got = wsv[c.coordinate].value
                check(expected is not None and abs(got - expected) < 1e-9,
                      f"{ws.title}!{c.coordinate} caches {got} but "
                      f"{code} {year} is {expected}")
                checked_values += 1

    # -- 7. charts are legible ------------------------------------------------
    charts = 0
    for ws in wb.worksheets:
        for ch in ws._charts:
            charts += 1
            where = f"{ws.title} chart {charts}"

            def text(obj):
                try:
                    return "".join(r.t or "" for p in obj.tx.rich.p
                                   for r in (p.r or []))
                except Exception:
                    return ""

            check(ch.title is not None and text(ch.title).strip(),
                  f"{where} has no title on the chart object")
            check(ch.x_axis.title is not None and text(ch.x_axis.title).strip()
                  or ch.x_axis.delete,
                  f"{where} has no x-axis title")
            check(ch.y_axis.title is not None and text(ch.y_axis.title).strip(),
                  f"{where} has no y-axis title")
            check(len(ch.series) > 0, f"{where} has no data series")

            # A scatter series must not join its points. XlsxWriter keeps any
            # truthy subtype string it is handed and only suppresses the
            # connecting line for the exact subtype "marker_only", so a typo
            # there had Excel drawing a zigzag through the 18 countries in row
            # order. The workbook has no scatter where joining points is
            # meaningful, so the rule is absolute.
            if type(ch).__name__ == "ScatterChart":
                for i, ser in enumerate(ch.series):
                    gp = ser.graphicalProperties
                    joined = not (gp is not None and gp.line is not None
                                  and gp.line.noFill)
                    check(not joined,
                          f"{where} series {i} joins its points with a line; "
                          f"a scatter of independent countries must not")

            # Value axes state their own range. Excel's autoscale put a 57-96%
            # series on a 0-120% axis, wasting half the plot.
            for name, ax in (("x", ch.x_axis), ("y", ch.y_axis)):
                if ax.delete or getattr(ax, "scaling", None) is None:
                    continue
                if type(ax).__name__ != "NumericAxis":
                    continue
                check(ax.scaling.min is not None and ax.scaling.max is not None,
                      f"{where} {name}-axis has no explicit min/max, so Excel "
                      f"will autoscale it")

            # -- 8. no chart sits on top of a populated cell ------------------
            r0, c0, r1, c1 = cell_extent(ch.anchor)
            collisions = []
            for rr in range(r0 + 1, r1 + 2):
                for cc in range(c0 + 1, c1 + 2):
                    if ws.cell(row=rr, column=cc).value not in (None, ""):
                        collisions.append(ws.cell(row=rr, column=cc).coordinate)
            check(not collisions,
                  f"{where} covers populated cells: {collisions[:8]}")
    check(charts == 4, f"expected 4 charts, found {charts}")

    # -- 9. every sheet has its furniture -------------------------------------
    for ws in wb.worksheets:
        check(len(ws.column_dimensions) > 0,
              f"{ws.title} has no column widths set")
        check(ws.sheet_properties.tabColor is not None,
              f"{ws.title} has no tab colour")
        check(ws.page_setup.orientation in ("portrait", "landscape"),
              f"{ws.title} has no print orientation")
        pr = ws.sheet_properties.pageSetUpPr
        check(pr is not None and pr.fitToPage,
              f"{ws.title} is not set to fit to page width when printed")

    # -- 10. the correction record survives -----------------------------------
    #    A correction that leaves no trace is indistinguishable from never
    #    having made the error, which is what the AI appendix must be able to
    #    show. Fail the build if the SMV:Digital withdrawal is ever dropped.
    lim = wb["07_LIMITATIONS"]
    blob = " ".join(str(c.value) for row in lim.iter_rows() for c in row
                    if c.value)
    check("SMV:Digital" in blob and "withdraw" in blob.lower(),
          "07_LIMITATIONS no longer records the withdrawn SMV:Digital claim")
    check("fabricated" in blob.lower(),
          "07_LIMITATIONS no longer records the discarded fabricated dataset")

    print(f"{CHECKS[0]} checks, {formulas} formulas "
          f"({checked_values} values cross-checked against the dataset), "
          f"{charts} charts, {len(wb.worksheets)} sheets")
    if FAILURES:
        print(f"\n{len(FAILURES)} FAILED:")
        for f in FAILURES:
            print("  -", f)
        return 1
    print("all checks pass")
    return 0


if __name__ == "__main__":
    sys.exit(main())

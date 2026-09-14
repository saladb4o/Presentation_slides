#!/usr/bin/env python3
"""Build the submission .docx from the shared content model.

Formatting follows the assessment brief: Times New Roman 12, 1.5 line spacing,
justified body, 1-inch margins, numbered sections, captioned figures.

Content comes from render.build(), the same call report/assemble_report.py
makes, so the Word file and the markdown preview carry identical prose,
citations, figure numbers and references.
"""
import re
import pathlib
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "report"))
from render import build, FIGURES_DIR                       # noqa: E402

BODY_FONT = "Times New Roman"
CLASS_GROUP = "[CLASS GROUP]"      # replaced once the code is known


def style_base(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(12)
    # Word needs the east-asian font set too, or it silently substitutes.
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    for section in doc.sections:
        # python-docx defaults to US Letter. An Australian university submission
        # is A4, and the difference shows on every page of a printed report.
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, Cm(2.54))


def _emphasis(text):
    """Split Markdown emphasis into (text, bold, italic) runs. Doubles first,
    so **x** is bold rather than italic-then-stray-asterisk."""
    out = []
    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            out.append((piece[2:-2], True, False))
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            out.append((piece[1:-1], False, True))
        else:
            out.append((piece, False, False))
    return out


def para(doc, text, *, size=12, bold=False, italic=False, align=None,
         space_before=0, space_after=6, color=None, spacing=None):
    p = doc.add_paragraph()
    # The drafts use Markdown emphasis, and a single run rendered it as literal
    # asterisks: every Danish statute name printed as *Lov om ...*. Split the
    # text into runs instead, so **bold** and *italic* become formatting.
    for piece, em_bold, em_italic in _emphasis(text):
        run = p.add_run(piece)
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.bold = bold or em_bold
        run.italic = italic or em_italic
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if spacing:
        p.paragraph_format.line_spacing_rule = spacing
    return p


def page_number_footer(doc):
    """Insert a PAGE field. python-docx has no API for this, so build the XML."""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def main():
    r = build()
    doc = Document()
    style_base(doc)
    page_number_footer(doc)

    # --- title block ------------------------------------------------------
    para(doc, r["meta"]["title"], size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    for line in r["meta"]["lines"]:
        para(doc, line.replace("[TO BE CONFIRMED]", CLASS_GROUP), size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "", size=6, space_after=0)
    thesis = para(doc, r["thesis"], size=12, italic=True, space_before=8,
                  space_after=14)
    thesis.paragraph_format.left_indent = Cm(1.0)
    thesis.paragraph_format.right_indent = Cm(1.0)

    for section in r["sections"]:
        para(doc, section["heading"], size=13, bold=True,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6)
        for text in section["paragraphs"]:
            para(doc, text)

        if section["key"] == "Q1_industries":
            t = r["table1"]
            cap = para(doc, t["caption"], size=10, bold=True,
                       align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8,
                       space_after=4)
            cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            # Width comes from the header, not a constant: a column added in
            # render.py was silently dropped when this was hardcoded to 3.
            table = doc.add_table(rows=1, cols=len(t["header"]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, head in zip(table.rows[0].cells, t["header"]):
                run = cell.paragraphs[0].add_run(head)
                run.bold, run.font.size, run.font.name = True, Pt(10), BODY_FONT
            for row in t["rows"]:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    run = cell.paragraphs[0].add_run(value)
                    run.font.size, run.font.name = Pt(10), BODY_FONT
            for trow in table.rows:
                for cell in trow.cells:
                    pf = cell.paragraphs[0].paragraph_format
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    pf.space_after = Pt(2)
            para(doc, "", size=6, space_after=0)

        for fig in section["figures"]:
            path = FIGURES_DIR / fig["file"]
            if not path.exists():
                raise SystemExit(f"missing figure image: {path} "
                                 "- run report/build_figures.py first")
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(8)
            pic.paragraph_format.space_after = Pt(2)
            pic.add_run().add_picture(str(path), width=Cm(15.5))
            cap = para(doc, fig["caption"], size=10, italic=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
            cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    doc.add_page_break()
    para(doc, "References", size=13, bold=True,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
    for ref in r["references"]:
        p = para(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
        # Harvard hanging indent.
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)

    for app in r["appendices"]:
        doc.add_page_break()
        para(doc, app["heading"], size=13, bold=True,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
        for kind, payload in app["blocks"]:
            if kind == "subheading":
                para(doc, payload, size=12, bold=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=4)
            elif kind == "para":
                para(doc, payload)
            elif kind == "bullets":
                for item in payload:
                    bp = doc.add_paragraph(style="List Bullet")
                    run = bp.add_run(item)
                    run.font.name, run.font.size = BODY_FONT, Pt(12)
                    bp.paragraph_format.space_after = Pt(3)
            elif kind == "table":
                tbl = doc.add_table(rows=0, cols=len(payload[0]))
                tbl.style = "Table Grid"
                for i, row in enumerate(payload):
                    cells = tbl.add_row().cells
                    for cell, value in zip(cells, row):
                        run = cell.paragraphs[0].add_run(value)
                        run.font.size, run.font.name = Pt(10), BODY_FONT
                        run.bold = (i == 0)
                        pf = cell.paragraphs[0].paragraph_format
                        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                        pf.space_after = Pt(2)
                para(doc, "", size=6, space_after=0)
        for fig in app["figures"]:
            path = FIGURES_DIR / fig["file"]
            if not path.exists():
                raise SystemExit(f"missing appendix figure: {path}")
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(8)
            pic.paragraph_format.space_after = Pt(2)
            pic.add_run().add_picture(str(path), width=Cm(15.5))
            cap = para(doc, fig["caption"], size=10, italic=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
            cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    name = f"Assignment 2 ECON1596 _ {CLASS_GROUP} _ s4040040.docx"
    out = ROOT / name
    doc.save(str(out))
    c = r["counts"]
    print(f"wrote {name}")
    print(f"  body {c['total']} words ({c['total_with_table']} if tables count), "
          f"{r['figure_count']} figures, {len(r['references'])} references")
    print(f"  {len(r['appendices'])} appendices: {r['appendix_words']} words, "
          f"{r['appendix_figure_count']} figures (outside the word count)")
    return out


if __name__ == "__main__":
    main()
